import urllib.request
import urllib.parse
import json
import io
import mimetypes

BASE_URL = "http://127.0.0.1:8000"

def test_api():
    print("Testing Root API...")
    req = urllib.request.Request(f"{BASE_URL}/")
    with urllib.request.urlopen(req) as resp:
        print("Root Status:", resp.status, resp.read().decode())

    print("\nTesting User Registration...")
    reg_data = json.dumps({
        "name": "Test User",
        "email": "testuser@gmail.com",
        "password": "password123",
        "title": "Software Engineer"
    }).encode("utf-8")
    
    try:
        req = urllib.request.Request(
            f"{BASE_URL}/auth/register",
            data=reg_data,
            headers={"Content-Type": "application/json"}
        )
        with urllib.request.urlopen(req) as resp:
            print("Register Status:", resp.status, resp.read().decode())
    except urllib.error.HTTPError as e:
        print("Register response (already exists or status):", e.code, e.read().decode())

    print("\nTesting User Login...")
    login_data = json.dumps({
        "email": "admin@gmail.com",
        "password": "123456"
    }).encode("utf-8")
    req = urllib.request.Request(
        f"{BASE_URL}/auth/login",
        data=login_data,
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        print("Login Status:", resp.status, resp.read().decode())

    print("\nTesting Resume Upload & Analysis...")
    boundary = "----WebKitFormBoundary7MA4YWxkTrZu0gW"
    
    # Create sample docx in memory
    from docx import Document
    doc = Document()
    doc.add_heading("John Doe - Software Engineer", level=1)
    doc.add_paragraph("Email: john.doe@gmail.com | Phone: 9876543210 | github.com/johndoe")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("Bachelor of Technology in Computer Science, 2024")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, JavaScript, React.js, FastAPI, Docker, SQL, Git, Linux, Machine Learning")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Developed and deployed scalable REST APIs with FastAPI and PostgreSQL. Optimized backend response time by 40%. Spearheaded automation test pipelines.")
    doc.add_heading("Projects", level=2)
    doc.add_paragraph("Engineered an AI Resume Analyzer using Python, NLP, and modern frontend tools.")
    
    bio = io.BytesIO()
    doc.save(bio)
    file_bytes = bio.getvalue()

    body = io.BytesIO()
    # file field
    body.write(f"--{boundary}\r\n".encode())
    body.write(b'Content-Disposition: form-data; name="file"; filename="sample_resume.docx"\r\n')
    body.write(b'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n')
    body.write(file_bytes)
    body.write(b"\r\n")
    # user_email field
    body.write(f"--{boundary}\r\n".encode())
    body.write(b'Content-Disposition: form-data; name="user_email"\r\n\r\n')
    body.write(b'admin@gmail.com\r\n')
    # target_role field
    body.write(f"--{boundary}\r\n".encode())
    body.write(b'Content-Disposition: form-data; name="target_role"\r\n\r\n')
    body.write(b'software_engineer\r\n')
    body.write(f"--{boundary}--\r\n".encode())

    req = urllib.request.Request(
        f"{BASE_URL}/analyze",
        data=body.getvalue(),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"}
    )
    with urllib.request.urlopen(req) as resp:
        res_json = json.loads(resp.read().decode())
        print("Analysis Status:", resp.status)
        print("ATS Score:", res_json.get("ats_score"))
        print("Skills Detected:", res_json.get("skills"))
        print("Missing Skills:", res_json.get("missing_skills"))
        print("Word Count:", res_json.get("word_count"))

    print("\nTesting Resume History Query...")
    req = urllib.request.Request(f"{BASE_URL}/history?email=admin@gmail.com")
    with urllib.request.urlopen(req) as resp:
        hist_json = json.loads(resp.read().decode())
        print("History Count:", len(hist_json.get("history", [])))
        print("Latest History Item ATS Score:", hist_json["history"][0]["ats_score"])

    print("\nAll Backend Tests Completed Successfully! [SUCCESS]")

if __name__ == "__main__":
    test_api()
